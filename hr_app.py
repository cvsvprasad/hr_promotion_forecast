# -*- coding: utf-8 -*-
"""
Created on Fri Feb  6 10:27:58 2026

@author: lenovo
"""
import streamlit as st
import pandas as pd
from datetime import datetime
from dateutil.relativedelta import relativedelta
from collections import defaultdict
import io, time, re
# ================= PDF =================
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib import colors

# ================= EXCEL CHARTS =================
from openpyxl.chart import BarChart, Reference

# =================================================
# LOGIN
# =================================================
def login():
    if "logged_in" not in st.session_state:
        st.session_state.logged_in = False
    if st.session_state.logged_in:
        return True

    st.title("🔐 HR Forecast Login")
    u = st.text_input("Username")
    p = st.text_input("Password", type="password")

    if st.button("Login"):
        if u == st.secrets["auth"]["username"] and p == st.secrets["auth"]["password"]:
            st.session_state.logged_in = True
            st.session_state.login_time = time.time()
            st.rerun()
        else:
            st.error("Invalid credentials")
    return False

if not login():
    st.stop()

if time.time() - st.session_state.login_time > 900:
    st.session_state.logged_in = False
    st.warning("Session expired")
    st.stop()

# =================================================
# PAGE
# =================================================
st.title("🏛 HR Promotion & Retirement Forecast System")
st.caption("For official planning use only. Data processed temporarily.")

# =================================================
# GOOGLE DRIVE (ROBUST)
# =================================================
DEFAULT_DRIVE_LINK = "https://docs.google.com/spreadsheets/d/1QMmRIjXJaaiaPqwW4qOxp2i2GPxwBJuQ/edit?usp=sharing&ouid=105054648927314914153&rtpof=true&sd=true"

def extract_drive_id(link):
    patterns = [r"/d/([a-zA-Z0-9_-]+)", r"id=([a-zA-Z0-9_-]+)"]
    for p in patterns:
        m = re.search(p, link)
        if m:
            return m.group(1)
    return None

def load_drive_excel(link):
    try:
        fid = extract_drive_id(link)
        if not fid:
            return None
        url = f"https://drive.google.com/uc?id={fid}"
        return pd.read_excel(url)
    except:
        return None

# =================================================
# INPUT SOURCE
# =================================================
st.subheader("Choose Data Source")
source = st.radio(
    "",
    ["Use Developer Default File", "Upload Excel File", "Load Custom Google Drive Link"]
)

df = None

if source == "Use Developer Default File":
    df = load_drive_excel(DEFAULT_DRIVE_LINK)
    if df is None:
        st.error("Default file not accessible (check Google Drive sharing)")
    else:
        st.success("Default file loaded")

elif source == "Upload Excel File":
    f = st.file_uploader("Upload Excel", type=["xlsx"])
    if f:
        df = pd.read_excel(f)

elif source == "Load Custom Google Drive Link":
    link = st.text_input("Paste Google Drive Excel Link")
    if link:
        df = load_drive_excel(link)
        if df is None:
            st.error("Invalid or private Google Drive file")

# =================================================
# UTILITIES
# =================================================
def parse_yy(d):
    d = datetime.strptime(str(d), "%d.%m.%y")
    if d.year > datetime.now().year:
        d = d.replace(year=d.year - 100)
    return d

def fmt(d):
    return d.strftime("%d-%m-%Y")

# =================================================
# PDF GENERATOR (A4, FULL)
# =================================================
def generate_pdf(master_df, promo_df, year_df):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="MyTitle", alignment=TA_CENTER, fontSize=14))

    story = []
    story.append(Paragraph("HR PROMOTION & RETIREMENT FORECAST REPORT", styles["MyTitle"]))
    story.append(Spacer(1, 12))

    def add_table(df):
        t = Table([df.columns.tolist()] + df.values.tolist(), repeatRows=1)
        t.setStyle(TableStyle([
            ("GRID",(0,0),(-1,-1),0.5,colors.black),
            ("BACKGROUND",(0,0),(-1,0),colors.lightgrey),
            ("FONTSIZE",(0,0),(-1,-1),8)
        ]))
        return t

    story.append(Paragraph("Master Data", styles["Heading2"]))
    story.append(add_table(master_df))
    story.append(PageBreak())

    story.append(Paragraph("Promotion History", styles["Heading2"]))
    story.append(add_table(promo_df))
    story.append(PageBreak())

    story.append(Paragraph("Year-wise Forecast", styles["Heading2"]))
    story.append(add_table(year_df))

    doc.build(story)
    buf.seek(0)
    return buf

# =================================================
# CORE PROCESS
# =================================================
if df is not None:

    df["DOB"] = df["DOB"].apply(parse_yy)
    df["Date of Retirement"] = df["DOB"].apply(lambda x: x + relativedelta(years=60))

    employees = []
    for i,r in df.iterrows():
        employees.append({
            "index": i,
            "sno": int(r["SNo"]),
            "name": r["Name Details"],
            "rank": int(r["Rank"]),
            "retire": r["Date of Retirement"],
            "retired": False,
            "history": []
        })

    retire_order = sorted(range(len(employees)), key=lambda i: employees[i]["retire"])

    promo_log = []
    yearly = defaultdict(lambda: {"ret":0,"pro":0})
    calendar = []
    promo_no = 1

    for idx in retire_order:
        emp = employees[idx]
        if emp["retired"]:
            continue

        emp["retired"] = True
        y = emp["retire"].year
        yearly[y]["ret"] += 1
        calendar.append(emp["retire"])

        vac = emp["rank"]
        while True:
            pool = [
                e for e in employees
                if not e["retired"]
                and e["rank"] == vac - 1
                and e["retire"] >= emp["retire"]
            ]
            if not pool:
                break

            p = sorted(pool, key=lambda x:x["sno"])[0]
            old,new = p["rank"], p["rank"] + 1
            p["rank"] = new
            p["history"].append(f"{old}→{new} on {fmt(emp['retire'])}")

            promo_log.append([promo_no, p["sno"], p["name"], old, new, fmt(emp["retire"])])
            yearly[y]["pro"] += 1

            promo_no += 1
            vac -= 1

    master_df = pd.DataFrame([
        [e["sno"], e["name"], fmt(df.loc[e["index"],"DOB"]),
         fmt(e["retire"]), " | ".join(e["history"]) if e["history"] else "—", e["rank"]]
        for e in sorted(employees, key=lambda x:x["sno"])
    ], columns=["SNo","Name","DOB","Retirement","Promotion History","Final Rank"])

    promo_df = pd.DataFrame(promo_log, columns=[
        "Promo No","SNo","Name","Old Rank","New Rank","Promotion Date"
    ])

    year_df = pd.DataFrame(
        [[y, yearly[y]["ret"], yearly[y]["pro"]] for y in sorted(yearly)],
        columns=["Year","Retirements","Promotions"]
    )

    # ================= DISPLAY =================
    st.subheader("📋 Master Data")
    st.dataframe(master_df, use_container_width=True)

    st.subheader("📊 Year-wise Forecast")
    st.bar_chart(year_df.set_index("Year"))

    # ================= EXCEL WITH CHARTS =================
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as w:
        master_df.to_excel(w, "Master_Data", index=False)
        promo_df.to_excel(w, "Promotion_History", index=False)
        year_df.to_excel(w, "Yearly_Forecast", index=False)

        ws = w.book["Yearly_Forecast"]
        chart = BarChart()
        chart.title = "Year-wise Retirements vs Promotions"
        data = Reference(ws, min_col=2, min_row=1, max_col=3, max_row=ws.max_row)
        cats = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, "E2")

    out.seek(0)
    st.download_button("⬇ Download Excel (with Charts)", out, "HR_Forecast.xlsx")

    # ================= PDF =================
    pdf_buf = generate_pdf(master_df, promo_df, year_df)
    st.download_button("🖨 Download PDF Report", pdf_buf, "HR_Forecast_Report.pdf")

# =================================================
# LOGOUT
# =================================================
st.markdown("---")
if st.button("Logout"):
    st.session_state.logged_in = False
    st.rerun()

